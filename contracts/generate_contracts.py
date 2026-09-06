#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_contracts.py — Dee Cleaners Househelp Agreements (EN + SW)
====================================================================
Generates the bilingual Dee Cleaners househelp agreement documents:

  Dee_Cleaners_Househelp_Agreement_EN.docx / .pdf
  Dee_Cleaners_Househelp_Agreement_SW.docx / .pdf

The terms are taken VERBATIM (translated faithfully into Swahili) from the
supplied "HOUSEHELP CLIENT FORM.pdf" — Dee Cleaners Househelp Terms &
Conditions. No additional contractual obligations are invented.

Requires: python-docx, reportlab
Run:      python3 contracts/generate_contracts.py
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image,
                                Table, TableStyle, HRFlowable)

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, '..', 'assets', 'dee_cleaners_logo', 'dee_cleaners_logo.png')

NAVY = RGBColor(0x21, 0x36, 0x68)
NAVY_HX = HexColor('#213668')
SKY_HX = HexColor('#2D9ED6')
GRAY_HX = HexColor('#5a6a85')

BUSINESS = {
    'en': {
        'name': 'Dee Cleaners',
        'address': 'Magomeni, Mikumi',
        'phones': '+255 693 806 051 / +255 762 292 439',
        'email': 'deecleanerscleaningservices@gmail.com',
        'title': 'Househelp Employment Agreement',
        'subtitle': 'Househelp Agreement Form & Terms and Conditions',
        'doc_id': 'Dee_Cleaners_Househelp_Agreement_EN',
        'parties': 'Parties',
        'employer_block': 'EMPLOYER / CLIENT INFORMATION',
        'househelp_block': 'HOUSEHELP INFORMATION',
        'fields': ['Full Name', 'Family Member / Relative', 'Phone Number',
                   'Email', 'Gender (Male/Female)', 'Nationality',
                   'Occupation', 'City/Country', 'NIDA Number',
                   'Home Address'],
        'terms_title': 'Terms and Conditions',
        'terms_intro': 'These Terms and Conditions apply to all parties involved '
                       '(Househelp and Employer) when engaging with Dee Cleaners. '
                       'By signing the Househelp Agreement Form, both parties agree '
                       'to the following terms:',
        'decl_emp_title': "Employer's Declaration",
        'decl_emp': 'I, …………………………………………, agree to adhere to the terms and '
                    'conditions listed above and treat the househelp with fairness '
                    'and respect.',
        'decl_hh_title': "Househelp's Declaration",
        'decl_hh': 'I, …………………………………………, agree to perform my duties diligently '
                   'and comply with the terms and conditions listed above.',
        'sign_note': 'By signing this form, both the employer and househelp confirm '
                     'their understanding and acceptance of the above terms.',
        'sig_employer': 'Signature (Employer)',
        'sig_househelp': 'Signature (Househelp)',
        'sig_date': 'Date',
        'footer': 'Dee Cleaners — Magomeni, Mikumi · +255 693 806 051 / +255 762 292 439 · deecleanerscleaningservices@gmail.com',
    },
    'sw': {
        'name': 'Dee Cleaners',
        'address': 'Magomeni, Mikumi',
        'phones': '+255 693 806 051 / +255 762 292 439',
        'email': 'deecleanerscleaningservices@gmail.com',
        'title': 'Mkataba wa Ajira ya Mfanyakazi wa Nyumbani',
        'subtitle': 'Fomu ya Makubaliano ya Mfanyakazi wa Nyumbani na Masharti',
        'doc_id': 'Dee_Cleaners_Househelp_Agreement_SW',
        'parties': 'Wahusika',
        'employer_block': 'TAARIFA ZA MWAJIRI / MTEJA',
        'househelp_block': 'TAARIFA ZA MFANYAKAZI WA NYUMBANI',
        'fields': ['Jina Kamili', 'Mwanafamilia / Jamaa', 'Namba ya Simu',
                   'Barua Pepe', 'Jinsia (Mwanamume/Mwanamke)', 'Uraia',
                   'Kazi', 'Mji/Nchi', 'Namba ya NIDA',
                   'Anwani ya Nyumbani'],
        'terms_title': 'Masharti na Vigezo',
        'terms_intro': 'Masharti haya yanawahusu wahusika wote (Mfanyakazi wa '
                       'Nyumbani na Mwajiri) wanaposhirikiana na Dee Cleaners. '
                       'Kwa kusaini Fomu ya Makubaliano ya Mfanyakazi wa Nyumbani, '
                       'pande zote mbili zinakubali masharti yafuatayo:',
        'decl_emp_title': 'Tamko la Mwajiri',
        'decl_emp': 'Mimi, …………………………………………, nakubali kuzingatia masharti '
                    'yaliyoorodheshwa hapo juu na kumtendea mfanyakazi wa nyumbani '
                    'kwa haki na heshima.',
        'decl_hh_title': 'Tamko la Mfanyakazi',
        'decl_hh': 'Mimi, …………………………………………, nakubali kutekeleza majukumu yangu '
                   'kwa bidii na kuzingatia masharti yaliyoorodheshwa hapo juu.',
        'sign_note': 'Kwa kusaini fomu hii, mwajiri na mfanyakazi wote '
                     'wanathibitisha uelewa na kukubali masharti yaliyo hapo juu.',
        'sig_employer': 'Sahihi (Mwajiri)',
        'sig_househelp': 'Sahihi (Mfanyakazi)',
        'sig_date': 'Tarehe',
        'footer': 'Dee Cleaners — Magomeni, Mikumi · +255 693 806 051 / +255 762 292 439 · deecleanerscleaningservices@gmail.com',
    },
}

TERMS = {
    'en': [
        ('1. Service Agreement', [
            'The househelp is employed to perform duties as outlined in the job description provided by the employer and agreed upon by Dee Cleaners.',
            'The employer must clearly specify the househelp’s work schedule, responsibilities, and expectations at the beginning of the employment.']),
        ('2. Working Hours', [
            'For Live-in Househelps: Standard working hours shall be from 5:00 a.m. to 6:30 p.m. with adequate meal breaks and rest periods during the day. The employer shall ensure that the househelp receives sufficient daily rest and is not required to work continuously without breaks.',
            'Any work beyond the agreed working hours shall be considered overtime and must be compensated as mutually agreed by both parties.']),
        ('3. Payment Terms', [
            'The employer agrees to pay the househelp on time, as stipulated in the signed agreement. The payment should be made monthly/weekly, by Cash, Mobile Payments or Bank deposits.',
            'Any salary deductions must be transparent and communicated in advance, e.g., for damages or non-performance.',
            'Dee Cleaners is entitled to a one-time agency fee paid by the employer upon contract initiation.']),
        ('4. Accommodation and Meals (for live-in househelps)', [
            'If the househelp is live-in, the employer must provide adequate accommodation (clean and safe living quarters) and sufficient meals or an allowance for food.',
            'Any additional expenses related to the househelp’s living conditions should be clearly agreed upon and documented.']),
        ('5. Workplace Environment', [
            'The employer agrees to provide a safe, respectful, and non-hostile working environment for the househelp.',
            'Any form of harassment, abuse, or discrimination by the employer or househelp will lead to immediate termination of the agreement.']),
        ('6. Termination of Employment', [
            'Either party may terminate this agreement by giving an oral or written notice and informing Dee Cleaners.',
            'If the employer delays or fails to pay the househelp’s salary on the agreed payment date, the househelp has the right to leave the job immediately and inform Dee Cleaners.',
            'In cases of abuse, harassment, unsafe working conditions, serious misconduct, or any situation that affects the househelp’s well-being, the househelp may terminate the agreement immediately and inform Dee Cleaners.',
            'In the event of immediate termination, Dee Cleaners must be notified within 24 hours.',
            'If the employer delays or fails to pay the househelp’s salary as agreed, Dee Cleaners reserves the right to terminate its services to the employer and may refuse to provide any replacement househelp or future placements.',
            'Dee Cleaners reserves the right to review each termination case individually and make a final decision based on the circumstances of the case.']),
        ('7. Health and Safety', [
            'The employer must provide a safe and hygienic environment for the househelp to work in.',
            'The househelp agrees to follow all reasonable health and safety procedures while on duty.']),
        ('8. Personal Conduct', [
            'The househelp is expected to maintain professionalism, integrity, and respect while carrying out duties.',
            'The employer agrees to treat the househelp with dignity and refrain from requesting duties outside of the agreed scope of work.']),
        ('9. Confidentiality', [
            'Both the employer and househelp agree to keep all personal and sensitive information obtained during employment confidential. This includes personal details of the employer’s family and the househelp.']),
        ('10. Dispute Resolution', [
            'Any disputes arising between the employer and the househelp will be mediated by Dee Cleaners.',
            'Both parties must cooperate with Dee Cleaners during mediation to resolve the issue amicably.']),
        ('11. Liability', [
            'Dee Cleaners is a third-party facilitator and cannot be held responsible for any disputes, theft, or damages that arise after the contract has been signed by both parties.',
            'The househelp is responsible for any intentional damage or loss of property while under employment, and the employer may deduct this from the salary after fair assessment.']),
        ('12. Mistreatment of Househelp', [
            'In cases where the househelp experiences mistreatment such as lack of food, inadequate clothing, physical abuse, or emotional abuse, the househelp reserves the right to terminate the contract immediately.',
            'Dee Cleaners must be notified of the mistreatment, and appropriate measures will be taken to resolve the issue.']),
        ('13. Minimum Work Duration', [
            'The househelp agrees to stay with the employer for a minimum of three (3) months if all agreed-upon terms and conditions are being met.',
            'Early termination by the househelp without valid reasons will result in replacement fees or additional penalties as stipulated in the agreement.']),
        ('14. Replacement Policy', [
            'If a Dee Cleaners househelp terminates the contract without valid reasons, Dee Cleaners will provide a replacement househelp to the employer at no additional cost.',
            'Replacement will be subject to availability and will follow the same agreed-upon terms.']),
        ('15. Payment Policy', [
            'All payments for the househelp’s services must be made directly to Dee Cleaners’ account.',
            'Dee Cleaners will handle the disbursement of salaries to the househelp.']),
        ('16. Customer Contract Breach', [
            'If a customer breaks the agreement, Dee Cleaners will not issue any refunds or provide another househelp.',
            'Dee Cleaners will only offer househelp services to customers who comply with the agreed-upon terms and conditions.']),
        ('17. Modifications to Terms', [
            'Any changes or amendments to the agreed terms must be communicated in writing and signed by both the employer and the househelp. Dee Cleaners must also be informed of any modifications.']),
    ],
    'sw': [
        ('1. Mkataba wa Huduma', [
            'Mfanyakazi wa nyumbani ameajiriwa kutekeleza majukumu kama yalivyoelezwa katika maelezo ya kazi yaliyotolewa na mwajiri na kukubaliwa na Dee Cleaners.',
            'Mwajiri lazima aeleze wazi ratiba ya kazi, majukumu, na matarajio ya mfanyakazi mwanzoni mwa ajira.']),
        ('2. Saa za Kazi', [
            'Kwa Wafanyakazi wa Ndani (Live-in): Saa za kawaida za kazi ni kuanzia 5:00 asubuhi hadi 6:30 jioni (saa 11:00 asubuhi – saa 12:30 jioni) pamoja na mapumziko ya kutosha ya chakula na kupumzika wakati wa siku. Mwajiri lazima ahakikishe mfanyakazi anapata mapumziko ya kutosha ya kila siku na hatakiwi kufanya kazi mfululizo bila mapumziko.',
            'Kazi yoyote zaidi ya saa za kazi zilizokubaliwa itachukuliwa kama kazi ya ziada (overtime) na lazima ilipwe kama walivyokubaliana pande zote mbili.']),
        ('3. Masharti ya Malipo', [
            'Mwajiri anakubali kumlipa mfanyakazi wa nyumbani kwa wakati, kama ilivyoelezwa katika mkataba uliotiwa saini. Malipo yafanyike kila mwezi/kila wiki, kwa Fedha Taslimu, Malipo ya Simu au Amana za Benki.',
            'Makato yoyote ya mshahara lazima yawe wazi na yatangazwe mapema, kwa mfano, kwa uharibifu au kushindwa kutekeleza majukumu.',
            'Dee Cleaners ana haki ya ada ya wakala ya mara moja inayolipwa na mwajiri mwanzoni mwa mkataba.']),
        ('4. Malazi na Milo (kwa wafanyakazi wa ndani)', [
            'Ikiwa mfanyakazi anaishi nyumbani (live-in), mwajiri lazima atoe malazi ya kutosha (makazi safi na salama) na milo ya kutosha au posho ya chakula.',
            'Gharama zozote za ziada zinazohusiana na hali ya maisha ya mfanyakazi zipaswa kukubaliwa wazi na kuandikwa.']),
        ('5. Mazingira ya Kazi', [
            'Mwajiri anakubali kutoa mazingira salama, yenye heshima, na yasiyo na uhasama kwa mfanyakazi wa nyumbani.',
            'Aina yoyote ya unyanyasaji, udhalilishaji, au ubaguzi kwa mwajiri au mfanyakazi itasababisha kusitishwa mara moja kwa mkataba.']),
        ('6. Kusitishwa kwa Ajira', [
            'Pande yoyote inaweza kusitisha mkataba huu kwa kutoa notisi ya mdomo au maandishi na kumjulisha Dee Cleaners.',
            'Ikiwa mwajiri anachelewesha au kushindwa kulipa mshahara wa mfanyakazi katika tarehe ya malipo iliyokubaliwa, mfanyakazi ana haki ya kuondoka mara moja na kumjulisha Dee Cleaners.',
            'Katika hali za unyanyasaji, usumbufu, mazingira hatari ya kazi, utovu mkubwa wa nidhamu, au hali yoyote inayoathiri ustawi wa mfanyakazi, mfanyakazi anaweza kusitisha mkataba mara moja na kumjulisha Dee Cleaners.',
            'Katika tukio la kusitishwa mara moja, Dee Cleaners lazima aarifiwe ndani ya saa 24.',
            'Ikiwa mwajiri anachelewesha au kushindwa kulipa mshahara wa mfanyakazi kama ilivyokubaliwa, Dee Cleaners ana haki ya kusitisha huduma zake kwa mwajiri na anaweza kukataa kutoa mfanyakazi mbadala au nafasi za baadaye.',
            'Dee Cleaners ana haki ya kupitia kila kesi ya kusitishwa kibinafsi na kufanya uamuzi wa mwisho kulingana na hali ya kesi.']),
        ('7. Afya na Usalama', [
            'Mwajiri lazima atoe mazingira salama na safi kwa mfanyakazi kufanyia kazi.',
            'Mfanyakazi anakubali kufuata taratibu zote za afya na usalama wakati wa kazi.']),
        ('8. Mwenendo Binafsi', [
            'Mfanyakazi anatarajiwa kudumisha taaluma, uadilifu, na heshima wakati wa kutekeleza majukumu.',
            'Mwajiri anakubali kumtendea mfanyakazi kwa hadhi na kujiepusha na kuomba majukumu nje ya upeo wa kazi uliokubaliwa.']),
        ('9. Usiri', [
            'Mwajiri na mfanyakazi wote wanakubali kuweka taarifa zote za kibinafsi na nyeti zilizopatikana wakati wa ajira kuwa siri. Hii inajumuisha maelezo ya kibinafsi ya familia ya mwajiri na ya mfanyakazi.']),
        ('10. Usuluhishi wa Migogoro', [
            'Migogoro yoyote inayojitokeza kati ya mwajiri na mfanyakazi itasuluhishwa na Dee Cleaners.',
            'Pande zote mbili lazima zishirikiane na Dee Cleaners wakati wa usuluhishi ili kutatua suala hilo kwa amani.']),
        ('11. Dhima (Liability)', [
            'Dee Cleaners ni mwezeshaji wa mtu wa tatu na hawezi kuwajibika kwa migogoro, wizi, au uharibifu wowote unaotokea baada ya mkataba kusainiwa na pande zote mbili.',
            'Mfanyakazi anawajibika kwa uharibifu wowote wa makusudi au upotevu wa mali wakati wa ajira, na mwajiri anaweza kukata hiyo kutoka kwenye mshahara baada ya tathmini ya haki.']),
        ('12. Unyanyasaji wa Mfanyakazi', [
            'Katika hali ambapo mfanyakazi anakumbana na unyanyasaji kama vile ukosefu wa chakula, mavazi yasiyotosha, unyanyasaji wa kimwili, au unyanyasaji wa kihisia, mfanyakazi ana haki ya kusitisha mkataba mara moja.',
            'Dee Cleaners lazima aarifiwe kuhusu unyanyasaji huo, na hatua stahiki zitachukuliwa kutatua suala hilo.']),
        ('13. Muda wa Chini wa Kazi', [
            'Mfanyakazi anakubali kukaa na mwajiri kwa muda usiopungua miezi mitatu (3) ikiwa masharti na vigezo vyote vilivyokubaliwa vinatimizwa.',
            'Kusitishwa mapema na mfanyakazi bila sababu za msingi kutasababisha ada za mbadala au adhabu za ziada kama ilivyoelezwa katika mkataba.']),
        ('14. Sera ya Mbadala', [
            'Ikiwa mfanyakazi wa Dee Cleaners atasitisha mkataba bila sababu za msingi, Dee Cleaners atampa mwajiri mfanyakazi mbadala bila gharama ya ziada.',
            'Mbadala atategemea upatikanaji na atafuata masharti yale yale yaliyokubaliwa.']),
        ('15. Sera ya Malipo', [
            'Malipo yote ya huduma za mfanyakazi lazima yafanywe moja kwa moja kwenye akaunti ya Dee Cleaners.',
            'Dee Cleaners atashughulikia ulipaji wa mishahara kwa mfanyakazi.']),
        ('16. Uvunjaji wa Mkataba na Mteja', [
            'Ikiwa mteja atavunja mkataba, Dee Cleaners hatarudisha fedha wala kutoa mfanyakazi mwingine.',
            'Dee Cleaners atatoa huduma za mfanyakazi kwa wateja wanaotii masharti yaliyokubaliwa pekee.']),
        ('17. Mabadiliko ya Masharti', [
            'Mabadiliko yoyote ya masharti yaliyokubaliwa lazima yawasilishwe kwa maandishi na kusainiwa na mwajiri na mfanyakazi. Dee Cleaners pia lazima aarifiwe kuhusu mabadiliko hayo.']),
    ],
}


# --------------------------------------------------------------------------
# DOCX generation
# --------------------------------------------------------------------------

def _field_row(doc, label, value=''):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{label}: ')
    r.bold = True
    p.add_run(value if value else '……………………………………………………………………………')


def build_docx(lang):
    t = BUSINESS[lang]
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    # header: logo + business name
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    cell_logo = table.cell(0, 0)
    cell_info = table.cell(0, 1)
    if os.path.exists(LOGO):
        cell_logo.paragraphs[0].add_run().add_picture(LOGO, width=Inches(1.15))
    p = cell_info.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(t['name'])
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = NAVY
    p2 = cell_info.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(t['subtitle'])
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x5A, 0x6A, 0x85)
    p3 = cell_info.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r3 = p3.add_run(f"{t['phones']}\n{t['email']}\n{t['address']}")
    r3.font.size = Pt(9)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = h.add_run(t['title'])
    hr.bold = True
    hr.font.size = Pt(15)
    hr.font.color.rgb = NAVY

    def section(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = NAVY

    # parties
    section(t['parties'])
    _field_row(doc, t['fields'][0] + ' (Mwajiri/Employer)' if lang == 'sw' else f"{t['fields'][0]} (Employer)")
    _field_row(doc, t['fields'][1])
    _field_row(doc, t['fields'][2])
    _field_row(doc, t['fields'][3])
    _field_row(doc, t['fields'][4])
    _field_row(doc, t['fields'][5])
    _field_row(doc, t['fields'][6])
    _field_row(doc, t['fields'][7])
    _field_row(doc, t['fields'][8])
    _field_row(doc, t['fields'][9])
    section(t['employer_block'] + ' — ' + t['fields'][0])
    _field_row(doc, t['fields'][0])
    doc.add_paragraph()

    # terms
    section(t['terms_title'])
    intro = doc.add_paragraph(t['terms_intro'])
    intro.paragraph_format.space_after = Pt(6)
    for title, items in TERMS[lang]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = NAVY
        for item in items:
            ip = doc.add_paragraph(item, style='List Bullet')
            ip.paragraph_format.space_after = Pt(1)

    # declarations
    section(t['decl_emp_title'])
    doc.add_paragraph(t['decl_emp'])
    section(t['decl_hh_title'])
    doc.add_paragraph(t['decl_hh'])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(t['sign_note']).italic = True

    # signatures
    sig = doc.add_table(rows=1, cols=3)
    for i, label in enumerate([t['sig_employer'], t['sig_househelp'], t['sig_date']]):
        c = sig.cell(0, i)
        cp = c.paragraphs[0]
        cp.paragraph_format.space_before = Pt(30)
        cp.add_run(label).bold = True
        cp.add_run('\n\n…………………………………')

    fp = doc.add_paragraph()
    fp.paragraph_format.space_before = Pt(18)
    fr = fp.add_run(t['footer'])
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x8B, 0x97, 0xAD)

    out = os.path.join(HERE, t['doc_id'] + '.docx')
    doc.save(out)
    return out


# --------------------------------------------------------------------------
# PDF generation (reportlab)
# --------------------------------------------------------------------------

def build_pdf(lang):
    t = BUSINESS[lang]
    out = os.path.join(HERE, t['doc_id'] + '.pdf')
    doc = SimpleDocTemplate(out, pagesize=A4,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            title=t['title'], author='Dee Cleaners')

    body = ParagraphStyle('body', fontName='Helvetica', fontSize=9.2, leading=12.6, spaceAfter=2)
    head = ParagraphStyle('head', fontName='Helvetica-Bold', fontSize=11.5,
                          leading=14, textColor=NAVY_HX, spaceBefore=8, spaceAfter=4)
    title_style = ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=15,
                                 leading=18, textColor=NAVY_HX, alignment=1, spaceAfter=2)
    sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=8.5, leading=11,
                         textColor=GRAY_HX, alignment=1, spaceAfter=10)

    story = []

    def esc(s):
        return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    logo_img = None
    if os.path.exists(LOGO):
        from reportlab.lib.utils import ImageReader
        logo_img = Image(LOGO, width=26 * mm, height=26 * mm)

    if logo_img:
        header = Table([[logo_img, Paragraph(
            f'<font size="16" color="#213668"><b>{esc(t["name"])}</b></font><br/>'
            f'<font size="8.5" color="#5a6a85">{esc(t["subtitle"])}<br/>'
            f'{esc(t["phones"])} · {esc(t["email"])}<br/>{esc(t["address"])}</font>',
            ParagraphStyle('ri', alignment=2, leading=11))]], colWidths=[32 * mm, 138 * mm])
        header.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ]))
        story.append(header)
        story.append(Spacer(1, 4 * mm))
        story.append(HRFlowable(width='100%', thickness=1, color=NAVY_HX))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(esc(t['title']), title_style))

    # parties fill-in table
    story.append(Paragraph(esc(t['parties']), head))
    fields = [f'{f}: ' + '_' * 46 for f in t['fields']]
    pt = Table([[Paragraph(esc(f), body)] for f in fields],
               colWidths=[170 * mm])
    pt.setStyle(TableStyle([
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
    ]))
    story.append(pt)

    # terms
    story.append(Paragraph(esc(t['terms_title']), head))
    story.append(Paragraph(esc(t['terms_intro']), body))
    for title, items in TERMS[lang]:
        story.append(Paragraph(esc(title), ParagraphStyle(
            'termhead', parent=head, fontSize=10, spaceBefore=6, spaceAfter=2)))
        for item in items:
            story.append(Paragraph('• ' + esc(item), body))

    # declarations
    story.append(Paragraph(esc(t['decl_emp_title']), head))
    story.append(Paragraph(esc(t['decl_emp']), body))
    story.append(Paragraph(esc(t['decl_hh_title']), head))
    story.append(Paragraph(esc(t['decl_hh']), body))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f'<i>{esc(t["sign_note"])}</i>', body))

    # signatures
    story.append(Spacer(1, 12 * mm))
    sig = Table([[Paragraph(esc(t['sig_employer']) + '<br/><br/>………………………………', ParagraphStyle('s', alignment=1)),
                  Paragraph(esc(t['sig_househelp']) + '<br/><br/>………………………………', ParagraphStyle('s', alignment=1)),
                  Paragraph(esc(t['sig_date']) + '<br/><br/>………………………………', ParagraphStyle('s', alignment=1))]],
                colWidths=[56.6 * mm, 56.6 * mm, 56.6 * mm])
    sig.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    story.append(sig)

    story.append(Spacer(1, 10 * mm))
    story.append(HRFlowable(width='100%', thickness=0.7, color=GRAY_HX))
    story.append(Paragraph(f'<font size="7.5" color="#8b97ad">{esc(t["footer"])}</font>',
                           ParagraphStyle('f', alignment=1, spaceBefore=3)))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(HexColor('#8b97ad'))
        canvas.drawString(16 * mm, 9 * mm, t['doc_id'])
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm,
                               f"{t['name']} — {t['address']}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out


if __name__ == '__main__':
    for lang in ('en', 'sw'):
        print('Generated', build_docx(lang))
        print('Generated', build_pdf(lang))
    print('All 4 contract documents ready in contracts/')
